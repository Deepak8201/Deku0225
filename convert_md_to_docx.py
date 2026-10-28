from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import sys

md_path = r"C:\Users\deku0225\Videos\AT task\cloud-jpg-tests\TestGeneration\tests\TestDesign\RTOSSFR-5350-OSPF-v2-delta.md"
docx_path = r"C:\Users\deku0225\Videos\AT task\cloud-jpg-tests\TestGeneration\tests\TestDesign\RTOSSFR-5350-OSPF-v2-delta.docx"

with open(md_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

in_table = False
table_lines = []

def add_formatted_text(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*|\`[^\`]*\`)', text)
    for part in parts:
        run = paragraph.add_run()
        if part.startswith('**') and part.endswith('**'):
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run.text = part[1:-1]
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        else:
            run.text = part

def flush_table():
    global in_table, table_lines
    if not table_lines:
        return
    rows = []
    for tl in table_lines:
        cells = [c.strip() for c in tl.strip().strip('|').split('|')]
        rows.append(cells)
    data_rows = [r for r in rows if not all(re.match(r'^[-:\s]+$', c) for c in r if c)]
    if not data_rows:
        in_table = False
        table_lines = []
        return
    max_cols = max(len(r) for r in data_rows)
    table = doc.add_table(rows=len(data_rows), cols=max_cols)
    table.style = 'Light Grid Accent 1'
    for i, row in enumerate(data_rows):
        for j in range(max_cols):
            cell_text = row[j] if j < len(row) else ''
            table.rows[i].cells[j].text = cell_text
    in_table = False
    table_lines = []

for raw_line in lines:
    line = raw_line.rstrip('\n')
    stripped = line.strip()

    if stripped == '---':
        flush_table()
        doc.add_paragraph()
        continue

    if not stripped:
        flush_table()
        continue

    if stripped.startswith('|'):
        in_table = True
        table_lines.append(line)
        continue
    else:
        if in_table:
            flush_table()

    if stripped.startswith('# '):
        p = doc.add_heading(stripped[2:], level=0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif stripped.startswith('## '):
        p = doc.add_heading(stripped[3:], level=1)
    elif stripped.startswith('### '):
        p = doc.add_heading(stripped[4:], level=2)
    elif stripped.startswith('#### '):
        p = doc.add_heading(stripped[5:], level=3)
    elif stripped.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        add_formatted_text(p, stripped[2:])
    elif re.match(r'^\d+\.\s+', stripped):
        p = doc.add_paragraph(style='List Number')
        add_formatted_text(p, re.sub(r'^\d+\.\s+', '', stripped))
    else:
        p = doc.add_paragraph()
        add_formatted_text(p, stripped)

flush_table()

doc.save(docx_path)
print(f"Saved: {docx_path}")
